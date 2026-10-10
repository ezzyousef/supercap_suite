"""Regression tests for the findings of the independent science review (docs/REVIEW_LOG.md)."""
import numpy as np
import pytest

from core import circuit_library as cl
from core import cv_analysis as cv
from core import eis_analysis as eis
from core import gcd_analysis as gcd


def _discharge(c_f=0.5, i_a=0.01, r_ohm=2.0, v_top=1.0, v_bottom=0.2, reversal=True):
    t = np.linspace(0, 200, 4001)
    step = (2 if reversal else 1) * i_a * r_ohm
    v = v_top - step - i_a * t / c_f
    t = np.concatenate([[0.0], t + 0.05])
    v = np.concatenate([[v_top], v])
    keep = v >= v_bottom
    return t[keep], v[keep]


def test_integral_capacitance_does_not_depend_on_where_the_window_ends():
    t = np.linspace(0, 100, 2001)
    for v_end in (0.0, 0.5, -1.0):              # e.g. a three-electrode window vs. a reference
        v = (v_end + 1.0) - t / 100.0           # 1 V linear discharge, C = I*dt/dV = 0.01*100/1
        assert gcd.total_capacitance_gcd_integral(t, v, 0.01) == pytest.approx(1.0, rel=1e-6)


def test_ir_drop_is_excluded_from_the_voltage_window():
    t, v = _discharge()
    r = gcd.total_capacitance_gcd_auto(t, v, 0.01)
    assert r["capacitance_f"] == pytest.approx(0.5, rel=1e-3)
    assert r["ir_drop_excluded_v"] == pytest.approx(0.04, rel=1e-3)
    kept = gcd.total_capacitance_gcd_auto(t, v, 0.01, exclude_ir_drop=False)
    assert kept["voltage_window_v"] > r["voltage_window_v"]


def test_esr_convention_for_a_current_reversal():
    t, v = _discharge(reversal=True)
    drop = gcd.estimate_ir_drop(t, v)
    assert gcd.esr_from_ir_drop(drop, 0.01, current_reverses=True) == pytest.approx(2.0, rel=1e-3)
    t, v = _discharge(reversal=False)
    drop = gcd.estimate_ir_drop(t, v)
    assert gcd.esr_from_ir_drop(drop, 0.01) == pytest.approx(2.0, rel=1e-3)


def test_cv_cycle_counter_sees_stacked_cycles():
    one = np.concatenate([np.linspace(0, 1, 100), np.linspace(1, 0, 100)[1:]])
    rng = np.random.default_rng(0)
    for n in (1, 2, 3):
        v = np.tile(one, n) + rng.normal(0, 0.003, one.size * n)
        assert cv.count_cycles(v) == pytest.approx(n, abs=0.25)


def test_classic_randles_is_in_the_library_and_recovers_its_parameters():
    spec = cl.get_circuit("randles_classic_C_W")
    f = np.logspace(5, -2, 60)
    true = {"Rs": 1.0, "Cdl": 2e-5, "Rct": 40.0, "Zw_Y0": 0.02}
    assert set(true) == set(spec.param_order)
    z = cl.evaluate_circuit(spec.tree, 2 * np.pi * f, true)
    fit = eis.fit_equivalent_circuit(f, z.real, z.imag, model="randles_classic_C_W", multistart=True)
    for k, val in true.items():
        assert fit.params[k] == pytest.approx(val, rel=0.02)


def test_a_large_device_capacitance_is_not_pinned_at_ten_farads():
    f = np.logspace(3, -2, 40)
    z = 0.02 + 1 / (1j * 2 * np.pi * f * 300.0)            # a 300 F cell
    fit = eis.fit_equivalent_circuit(f, z.real, z.imag, model="baseline_RC")
    assert fit.params["C"] == pytest.approx(300.0, rel=0.01)
    assert not any("upper search bound" in w for w in fit.warnings)


def test_modulus_weighting_reports_a_unit_free_residual():
    spec = cl.get_circuit("randles1_C_none")
    f = np.logspace(5, -1, 40)
    true = {"Rs": 2.0, "Rct": 50.0, "Rct_cap": 1e-5}
    z = cl.evaluate_circuit(spec.tree, 2 * np.pi * f, true)
    rng = np.random.default_rng(0)
    zn = z * (1 + 0.01 * rng.standard_normal(z.size))
    a = eis.fit_equivalent_circuit(f, zn.real, zn.imag, model="randles1_C_none")
    b = eis.fit_equivalent_circuit(f, zn.real * 1000, zn.imag * 1000, model="randles1_C_none")
    assert a.rel_rms_percent == pytest.approx(b.rel_rms_percent, rel=1e-3)   # mΩ vs kΩ: same
    assert a.rel_rms_percent < 2.0
    assert np.isfinite(a.aicc)


def test_kramers_kronig_accepts_a_valid_supercapacitor_spectrum():
    f = np.logspace(5, -2, 60)
    w = 2 * np.pi * f
    z = 0.5 + 1 / (1 / 2.0 + 1j * w * 1e-3) + 1 / (1j * w * 0.5)   # with a capacitive tail
    rng = np.random.default_rng(1)
    zn = z * (1 + 0.002 * rng.standard_normal(f.size))
    r = eis.kramers_kronig_test(f, zn.real, zn.imag)
    assert r.passed and r.max_residual_percent < 1.0
    assert r.series_capacitance_f == pytest.approx(0.5, rel=0.02)


def test_appending_rows_keeps_the_header_under_a_source_note(tmp_path):
    import pandas as pd
    from core import export_io

    path = tmp_path / "log.xlsx"
    first = pd.DataFrame({"File": ["a"], "C (F/g)": [100.0]})
    sheet = export_io.export_table(str(path), "Log", first, source_note="GCD — test suite")
    export_io.append_table_rows(str(path), sheet, pd.DataFrame({"File": ["b"], "C (F/g)": [110.0]}))
    back = pd.read_excel(path, sheet_name=sheet, header=2)
    assert list(back.columns) == ["File", "C (F/g)"]
    assert back["File"].tolist() == ["a", "b"]
    note = pd.read_excel(path, sheet_name=sheet, header=None).iloc[0, 0]
    assert str(note).startswith("GCD — test suite")


def test_dsc_heat_flow_units():
    from core import dsc_analysis as dsc
    assert dsc.guess_heat_flow_unit("Heat Flow (W/g)") == "W/g (normalised)"
    assert dsc.guess_heat_flow_unit("Heat flow/mW") == "mW"
    assert dsc.heat_flow_to_mw([1.0], "W")[0] == pytest.approx(1000.0)
    assert dsc.heat_flow_to_mw([0.5], "W/g (normalised)", 0.010)[0] == pytest.approx(5.0)
    with pytest.raises(ValueError):
        dsc.heat_flow_to_mw([0.5], "W/g (normalised)", None)


def test_european_semicolon_csv_reads_numbers(tmp_path):
    from core import data_io
    path = tmp_path / "eu.csv"
    path.write_text("time/s;Ewe/V;<I>/mA\n0,0;0,10;1,5\n1,0;0,20;1,5\n\n2,0;0,30;1,5\n", encoding="utf-8")
    df = data_io.load_data_file(str(path))
    assert list(df.columns) == ["time/s", "Ewe/V", "<I>/mA"]
    assert df["Ewe/V"].tolist() == pytest.approx([0.1, 0.2, 0.3])
    assert len(df) == 3


def test_coulombic_efficiency_uses_charge_when_current_is_given():
    # charge at 2 mA for 50 s, discharge at 1 mA for 90 s: CE by charge = 90 %,
    # whereas the time ratio would claim 180 %.
    t1 = np.linspace(0, 50, 501)
    t2 = 50 + np.linspace(0, 90, 901)[1:]
    t = np.concatenate([t1, t2] * 1)
    v = np.concatenate([0.02 * t1, 1.0 - (t2 - 50) / 90.0])
    i = np.concatenate([np.full(t1.size, 0.002), np.full(t2.size, -0.001)])
    t = np.concatenate([t, t[-1] + t[1:] - t[0] + 0.1])
    v = np.concatenate([v, v[1:]])
    i = np.concatenate([i, i[1:]])
    res = gcd.analyze_cycling_stability(t, v, 0.001, 0.005, current_series_a=i)
    assert res[0].coulombic_efficiency_percent == pytest.approx(90.0, rel=0.03)



def test_an_ideal_rectangle_scores_one():
    v = np.concatenate([np.linspace(0, 1, 100), np.linspace(1, 0, 100)])
    i = np.concatenate([np.full(100, 1e-3), np.full(100, -1e-3)])
    assert cv.assess_cv_rectangularity(v, i) == pytest.approx(1.0)


def test_reim_lambda_selection_recovers_two_zarc_peaks():
    from core import drt_analysis as drt
    f = np.logspace(5, -2, 60)
    w = 2 * np.pi * f
    z = 2 + 50 / (1 + (1j * w * 1e-3) ** 0.8) + 20 / (1 + (1j * w * 1e-1) ** 0.9)
    zn = z * (1 + 0.005 * np.random.default_rng(0).standard_normal(f.size))
    lam, _c, _s = drt.select_lambda_reim(f, zn.real, zn.imag)
    r = drt.compute_drt(f, zn.real, zn.imag, lambda_reg=lam)
    taus = sorted(p.tau_s for p in r.peaks if p.within_measured_range)
    assert len(taus) == 2
    assert taus[0] == pytest.approx(1e-3, rel=0.3) and taus[1] == pytest.approx(0.1, rel=0.3)


def test_second_export_into_an_existing_workbook_keeps_note_and_table_together(tmp_path):
    from openpyxl import load_workbook
    from core import export_io
    path = str(tmp_path / "x.xlsx")
    export_io.export_new_sheet(path, "GCD", {"C": 1.0}, None, "note A")
    export_io.export_new_sheet(path, "GCD", {"C": 2.0}, None, "note B")
    sheets = load_workbook(path).worksheets
    assert len(sheets) == 2
    for ws, note in zip(sheets, ("note A", "note B")):
        rows = [[c.value for c in r] for r in ws.iter_rows()]
        assert str(rows[0][0]).startswith(note) and rows[2] == ["Parameter", "Value"]


def test_overlapping_dsc_peaks_are_split_at_the_valley():
    from core import dsc_analysis as dsc
    t = np.linspace(0, 600, 3001)
    y = -(5 * np.exp(-0.5 * ((t - 250) / 20) ** 2) + 3 * np.exp(-0.5 * ((t - 330) / 20) ** 2))
    comps = dsc.split_overlapping_peaks(t, y, np.zeros_like(t))
    assert len(comps) == 2
    true = np.array([5, 3]) * 20 * np.sqrt(2 * np.pi) / 1000
    assert [c.area_j for c in comps] == pytest.approx(true, rel=0.03)
    single = dsc.split_overlapping_peaks(t, -5 * np.exp(-0.5 * ((t - 250) / 20) ** 2), np.zeros_like(t))
    assert len(single) == 1 and single[0].fraction == pytest.approx(1.0)


def test_kramers_kronig_flags_smooth_drift_that_still_passes():
    f = np.logspace(5, -2, 60)
    w = 2 * np.pi * f
    z = 0.5 + 1 / (1 / 2.0 + 1j * w * 1e-3) + 1 / (1j * w * 0.5)
    zn = z * (1 + 0.002 * np.random.default_rng(0).standard_normal(f.size))
    valid = eis.kramers_kronig_test(f, zn.real, zn.imag)
    assert valid.passed and not valid.systematic_trend
    zd = zn * (1 + 0.3 * np.linspace(0, 1, f.size))
    drift = eis.kramers_kronig_test(f, zd.real, zd.imag)
    assert drift.passed and drift.systematic_trend


def test_mw_per_gram_is_not_read_as_w_per_gram():
    from core import dsc_analysis as dsc
    assert dsc.guess_heat_flow_unit("Heat flow (mW/g)") == "mW/g (normalised)"
    assert dsc.heat_flow_to_mw([1.0], "mW/g (normalised)", 0.01)[0] == pytest.approx(0.01)


def test_ir_step_logged_at_the_same_timestamp_is_excluded():
    t = np.array([0, 0, 1, 2, 3, 4, 5, 6, 7, 8.0])
    v = np.array([1, .8, .7, .6, .5, .4, .3, .2, .1, 0])
    r = gcd.capacitance_gcd_auto(t, v, 1.0, 1.0)
    assert r["capacitance_f_per_g"] == pytest.approx(10.0)
    assert r["ir_drop_excluded_v"] == pytest.approx(0.2)
